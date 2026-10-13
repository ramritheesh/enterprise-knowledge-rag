
from pathlib import Path
import logging
import re
from io import BytesIO

import requests
from docx import Document
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from pypdf import PdfReader

from app.generation.answer import answer_question
from app.ingestion.vector_store import (
    collection,
    delete_document_chunks,
    index_documents,
)


logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = FastAPI(
    title="Enterprise Knowledge Intelligence Platform",
    description="A RAG API for answering questions from company documents.",
    version="1.6.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
    ],
    allow_credentials=False,
    allow_methods=["GET", "POST", "DELETE"],
    allow_headers=["Content-Type"],
)

DOCUMENTS_DIR = Path("data/documents")
MAX_FILE_SIZE = 2 * 1024 * 1024
ALLOWED_EXTENSIONS = {".txt", ".pdf", ".docx"}


class AskRequest(BaseModel):
    question: str = Field(..., min_length=1)
    top_k: int = Field(default=3, ge=1, le=5)


class Source(BaseModel):
    chunk_id: str | None = None
    source: str
    text: str
    distance: float


class AskResponse(BaseModel):
    answer: str
    sources: list[Source]


@app.get("/health")
def health_check():
    return {
        "status": "ok",
        "service": "enterprise-knowledge-rag",
    }


@app.post("/ask", response_model=AskResponse)
def ask_question(request: AskRequest):
    if not request.question.strip():
        raise HTTPException(
            status_code=422,
            detail="Question cannot be empty.",
        )

    try:
        return answer_question(
            question=request.question.strip(),
            top_k=request.top_k,
        )
    except ValueError as error:
        raise HTTPException(
            status_code=400,
            detail=str(error),
        ) from error
    except requests.exceptions.RequestException as error:
        raise HTTPException(
            status_code=502,
            detail=(
                "The local language model service is unavailable. "
                "Check that Ollama is running."
            ),
        ) from error


@app.get("/documents")
def list_documents():
    """List supported files and whether their names are indexed."""

    DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)

    indexed_data = collection.get(include=["metadatas"])

    indexed_sources = {
        (metadata or {}).get("source")
        for metadata in indexed_data.get("metadatas", [])
        if (metadata or {}).get("source")
    }

    documents = []

    for path in sorted(DOCUMENTS_DIR.iterdir()):
        if not path.is_file():
            continue

        if path.suffix.lower() not in ALLOWED_EXTENSIONS:
            continue

        documents.append(
            {
                "filename": path.name,
                "extension": path.suffix.lower(),
                "size_bytes": path.stat().st_size,
                "indexed": path.name in indexed_sources,
            }
        )

    return {
        "total_documents": len(documents),
        "indexed_documents": sum(
            1 for document in documents if document["indexed"]
        ),
        "documents": documents,
    }


def extract_uploaded_text(filename: str, content: bytes) -> str:
    """Extract text from supported TXT, PDF, and DOCX files."""

    extension = Path(filename).suffix.lower()

    try:
        if extension == ".txt":
            extracted_text = content.decode("utf-8")

        elif extension == ".pdf":
            reader = PdfReader(BytesIO(content))

            if reader.is_encrypted:
                raise HTTPException(
                    status_code=400,
                    detail="Password-protected PDFs are not supported.",
                )

            extracted_text = "\n\n".join(
                page.extract_text() or ""
                for page in reader.pages
            )

        elif extension == ".docx":
            document = Document(BytesIO(content))

            parts = [
                paragraph.text.strip()
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))

            extracted_text = "\n\n".join(parts)

        else:
            raise HTTPException(
                status_code=400,
                detail="Only .txt, .pdf, and .docx files are supported.",
            )

    except HTTPException:
        raise
    except UnicodeDecodeError as error:
        raise HTTPException(
            status_code=400,
            detail="The TXT file must use UTF-8 encoding.",
        ) from error
    except Exception as error:
        logger.warning(
            "Document text extraction failed for %s",
            Path(filename).name,
            exc_info=True,
        )
        raise HTTPException(
            status_code=400,
            detail=(
                "Could not read the document. Ensure it is a valid, "
                "unprotected PDF or DOCX file."
            ),
        ) from error

    if not extracted_text.strip():
        raise HTTPException(
            status_code=400,
            detail=(
                "No readable text was found. Scanned PDFs may require OCR."
            ),
        )

    return extracted_text


@app.post("/upload")
async def upload_document(file: UploadFile = File(...)):
    original_name = Path(file.filename or "").name

    if not original_name:
        raise HTTPException(
            status_code=400,
            detail="A filename is required.",
        )

    extension = Path(original_name).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="Only .txt, .pdf, and .docx files are supported.",
        )

    safe_name = re.sub(
        r"[^A-Za-z0-9._-]",
        "_",
        original_name,
    )

    if safe_name in {"", ".", ".."}:
        raise HTTPException(
            status_code=400,
            detail="Invalid filename.",
        )

    content = await file.read(MAX_FILE_SIZE + 1)
    await file.close()

    if not content:
        raise HTTPException(
            status_code=400,
            detail="The file is empty.",
        )

    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail="File exceeds the 2 MB upload limit.",
        )

    extracted_text = extract_uploaded_text(safe_name, content)

    DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)
    destination = DOCUMENTS_DIR / safe_name

    if destination.exists():
        raise HTTPException(
            status_code=409,
            detail="A file with this name already exists.",
        )

    try:
        destination.write_bytes(content)
        logger.info("Saved uploaded document: %s", safe_name)

        index_documents()

        # Confirm that this file has at least one indexed chunk.
        indexed = collection.get(
            where={"source": safe_name},
            include=["metadatas"],
        )

        if not indexed.get("ids"):
            raise RuntimeError(
                f"No indexed chunks found for {safe_name}"
            )

        logger.info("Successfully indexed %s", safe_name)

    except Exception as error:
        logger.exception("Upload/indexing failed for %s", safe_name)

        # Do not remove the source file if vector cleanup fails.
        cleanup_succeeded = False

        try:
            delete_document_chunks(safe_name)
            cleanup_succeeded = True
        except Exception:
            logger.exception(
                "Could not clean up indexed chunks for %s. "
                "The source file will be retained for recovery.",
                safe_name,
            )

        if cleanup_succeeded:
            try:
                destination.unlink(missing_ok=True)
            except OSError:
                logger.exception(
                    "Could not remove failed upload file %s",
                    safe_name,
                )

        if cleanup_succeeded:
            detail = (
                "Document indexing failed. Cleanup was attempted. "
                "Check the backend logs for details."
            )
        else:
            detail = (
                "Document indexing failed and vector cleanup also failed. "
                "The source file was retained for recovery. Check backend "
                "logs and reconcile the vector index before retrying."
            )

        raise HTTPException(
            status_code=500,
            detail=detail,
        ) from error

    return {
        "message": "Document uploaded and indexed successfully.",
        "filename": safe_name,
        "characters": len(extracted_text),
    }


@app.delete("/documents/{filename}")
def delete_document(filename: str):
    """Delete one stored document and only its indexed chunks."""

    if (
        not filename
        or filename in {".", ".."}
        or Path(filename).name != filename
        or "/" in filename
        or "\\" in filename
    ):
        raise HTTPException(
            status_code=400,
            detail="Invalid filename.",
        )

    if Path(filename).suffix.lower() not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="Only .txt, .pdf, and .docx documents can be deleted.",
        )

    destination = DOCUMENTS_DIR / filename

    if not destination.is_file():
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    try:
        original_content = destination.read_bytes()
    except OSError as error:
        logger.exception("Could not read %s before deletion", filename)
        raise HTTPException(
            status_code=500,
            detail="Could not prepare the document for deletion.",
        ) from error

    try:
        # Remove the selected document's vectors first.
        deleted_chunks = delete_document_chunks(filename)

        # Then remove its file from disk.
        destination.unlink()

        logger.info(
            "Deleted document %s and %d indexed chunks",
            filename,
            deleted_chunks,
        )

    except Exception as error:
        logger.exception("Deletion failed for %s", filename)

        # Restore the file if it may have been removed before failure.
        try:
            if not destination.exists():
                destination.write_bytes(original_content)

            index_documents()

        except Exception:
            logger.exception(
                "Recovery failed after deleting %s. "
                "The file and vector index may require manual reconciliation.",
                filename,
            )

        raise HTTPException(
            status_code=500,
            detail=(
                "Document deletion failed. Check the backend terminal "
                "and verify the document list and index."
            ),
        ) from error

    return {
        "message": "Document deleted successfully.",
        "filename": filename,
        "deleted_chunks": deleted_chunks,
    }
