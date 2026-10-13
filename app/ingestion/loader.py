
from pathlib import Path

from docx import Document
from pypdf import PdfReader


DOCUMENTS_DIR = Path("data/documents")

SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}


def extract_text(file_path: Path) -> str:
    """Extract text from a supported document."""

    extension = file_path.suffix.lower()

    if extension == ".txt":
        return file_path.read_text(encoding="utf-8")

    if extension == ".pdf":
        reader = PdfReader(str(file_path))

        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")

        return "\n\n".join(pages)

    if extension == ".docx":
        document = Document(str(file_path))

        paragraphs = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        # Also extract text from tables, which are common in handbooks.
        table_text = []
        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]
                if cells:
                    table_text.append(" | ".join(cells))

        return "\n\n".join(paragraphs + table_text)

    raise ValueError(
        f"Unsupported file type: {extension}. "
        "Supported types are TXT, PDF, and DOCX."
    )


def load_text_documents(
    directory: str | Path = DOCUMENTS_DIR,
) -> list[dict]:
    """Load supported documents and return their text and source names."""

    root = Path(directory)

    if not root.exists():
        return []

    documents = []

    for file_path in sorted(root.iterdir()):
        if not file_path.is_file():
            continue

        if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        try:
            text = extract_text(file_path)
        except Exception as error:
            print(f"Could not read {file_path.name}: {error}")
            continue

        if not text.strip():
            print(
                f"Warning: No extractable text found in {file_path.name}. "
                "Scanned PDFs may require OCR."
            )
            continue

        documents.append(
            {
                "text": text,
                "source": file_path.name,
            }
        )

    return documents
